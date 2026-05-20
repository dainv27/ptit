#!/usr/bin/env python3
"""Build BTL report DOCX from markdown (font 13, ~15 pages)."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(13)
LINE_SPACING = 1.15
MERMAID_DIAGRAMS = {
    "usecase": (
        "Hình 2.1 – Use case (tóm tắt):\n"
        "  Người quản lý → Kết nối Server, Quản lý khách sạn/phòng, Tìm kiếm, Xem danh sách\n"
        "  Người vận hành → Khởi động/Dừng Server\n"
        "  Các use case nghiệp vụ → Lưu/truy vấn MySQL"
    ),
    "arch": (
        "Hình 3.1 – Kiến trúc:\n"
        "  [Người dùng] → Client JavaFX (ClientController) → ServerConnection (TCP)\n"
        "       ↔ HotelTcpServer (ServerSocket) → HotelService → DatabaseManager → MySQL\n"
        "  [Người vận hành] → Server JavaFX (ServerController) → HotelTcpServer"
    ),
    "json": (
        "Hình 3.2 – Thông điệp JSON:\n"
        "  Request: { command, data }  →  Response: { ok, message, data }"
    ),
    "class": (
        "Hình 4.1 – Biểu đồ lớp (các lớp chính):\n"
        "  Client: ClientApplication → ClientController → ServerConnection → ServerResponse\n"
        "  Server: ServerApplication → ServerController → HotelTcpServer → HotelService\n"
        "  HotelService → DatabaseManager → DatabaseConfig; Hotel, Room (entity)"
    ),
    "seq_connect": (
        "Hình 4.2 – Tuần tự kết nối & LIST_HOTELS:\n"
        "  User → ClientController → ServerConnection.connect → HotelTcpServer\n"
        "  → send LIST_HOTELS → HotelService.listHotels → MySQL SELECT → trả JSON → cập nhật bảng"
    ),
    "seq_add_hotel": (
        "Hình 4.3 – Tuần tự ADD_HOTEL:\n"
        "  User nhập form → send ADD_HOTEL → HotelService.addHotel → INSERT hotels\n"
        "  → ok/message → LIST_HOTELS làm mới bảng (hoặc báo lỗi trùng mã)"
    ),
    "seq_search": (
        "Hình 4.4 – Tuần tự SEARCH_ROOMS:\n"
        "  User nhập loại/giá → SEARCH_ROOMS → searchRooms (SQL LIKE, price<=) → hiển thị kết quả"
    ),
    "er": (
        "Hình 4.5 – ER:\n"
        "  HOTELS (id PK, name, stars, description) 1 —— n ROOMS\n"
        "  ROOMS (hotel_id PK,FK, room_id PK, type, price)"
    ),
    "deploy": (
        "Hình 5.1 – Triển khai:\n"
        "  Máy Client (JavaFX) ←TCP 5555 JSON→ Máy Server (JavaFX + HotelTcpServer) ←JDBC→ MySQL"
    ),
}


def set_run_font(run, bold=False, italic=False, size=None):
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = size or FONT_SIZE
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_paragraph_format(paragraph, align=None, space_before=0, space_after=6, indent_cm=0):
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if indent_cm:
        pf.left_indent = Cm(indent_cm)
    if align is not None:
        paragraph.alignment = align


def add_paragraph(doc, text="", style=None, bold=False, align=None, space_after=6):
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, align=align, space_after=space_after)
    if text:
        run = p.add_run(text)
        set_run_font(run, bold=bold)
    return p


def add_heading(doc, text, level=1):
    style = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, space_before=12 if level == 1 else 8, space_after=6)
    run = p.add_run(text)
    set_run_font(run, bold=True, size=Pt(14 if level == 1 else 13))
    return p


def add_table_from_rows(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.rows[i].cells[j]
            cell.text = row[j] if j < len(row) else ""
            for paragraph in cell.paragraphs:
                set_paragraph_format(paragraph, space_after=0)
                for run in paragraph.runs:
                    set_run_font(run, bold=(i == 0))
    doc.add_paragraph()


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    set_paragraph_format(p, space_after=6)
    pf = p.paragraph_format
    pf.left_indent = Cm(1)
    text = "\n".join(lines)
    run = p.add_run(text)
    set_run_font(run, size=Pt(11))
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")


def describe_mermaid(code_lines, lang=""):
    joined = lang + " " + " ".join(code_lines)
    if "Manager" in joined and "UCConnect" in joined:
        return MERMAID_DIAGRAMS["usecase"]
    if "ClientUI" in joined or "HotelTcpServer" in joined and "Conn" in joined:
        return MERMAID_DIAGRAMS["arch"]
    if "Request JSON" in joined or "command: tên lệnh" in joined:
        return MERMAID_DIAGRAMS["json"]
    if "classDiagram" in joined or "ClientApplication" in joined:
        return MERMAID_DIAGRAMS["class"]
    if "ADD_HOTEL" in joined:
        return MERMAID_DIAGRAMS["seq_add_hotel"]
    if "SEARCH_ROOMS" in joined:
        return MERMAID_DIAGRAMS["seq_search"]
    if "LIST_HOTELS" in joined and "sequenceDiagram" in joined:
        return MERMAID_DIAGRAMS["seq_connect"]
    if "erDiagram" in joined or "HOTELS ||" in joined:
        return MERMAID_DIAGRAMS["er"]
    if "ClientMachine" in joined:
        return MERMAID_DIAGRAMS["deploy"]
    if "sequenceDiagram" in joined:
        return MERMAID_DIAGRAMS["seq_connect"]
    if "flowchart" in joined:
        return MERMAID_DIAGRAMS["arch"]
    return "Sơ đồ minh họa (tham chiếu mã nguồn client/server)."


def parse_table_block(lines):
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        if re.match(r"^\|[-:\s|]+\|$", line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
    return rows


def setup_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = FONT_SIZE
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    return doc


def add_cover(doc):
    for _ in range(3):
        add_paragraph(doc, "")
    add_paragraph(doc, "BÁO CÁO BÀI TẬP LỚN", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    add_paragraph(doc, "MÔN: LẬP TRÌNH MẠNG", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
    add_paragraph(doc, "Đề tài: Quản lý khách sạn bằng giao thức TCP/IP", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=36)
    for label, value in [
        ("Họ tên:", "Nguyễn Văn Đại"),
        ("MSSV:", "K24DTCN0101"),
        ("Lớp:", "D24TXCN01-K"),
        ("Hình thức:", "Báo cáo cá nhân"),
    ]:
        p = add_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
        r1 = p.add_run(f"{label} ")
        set_run_font(r1, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2)
    doc.add_page_break()


def convert_markdown(md_path: Path, out_path: Path):
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = setup_document()
    add_cover(doc)
    add_manual_toc(doc)

    i = 0
    in_code = False
    code_buf: list[str] = []
    code_lang = ""
    table_buf: list[str] = []
    skip_title_block = True

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if skip_title_block:
            if stripped.startswith("# Danh mục từ viết tắt"):
                skip_title_block = False
            else:
                i += 1
                continue

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = stripped[3:].strip()
                code_buf = []
            else:
                in_code = False
                if code_lang == "mermaid":
                    desc = describe_mermaid(code_buf, code_lang)
                    p = add_paragraph(doc)
                    run = p.add_run(desc)
                    set_run_font(run, italic=True)
                else:
                    add_code_block(doc, code_buf)
                code_buf = []
                code_lang = ""
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if stripped.startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        elif table_buf:
            add_table_from_rows(doc, parse_table_block(table_buf))
            table_buf = []

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("# ") and not stripped.startswith("## "):
            title = stripped[2:].strip()
            if title.upper() in ("KẾT LUẬN", "TÀI LIỆU THAM KHẢO"):
                add_heading(doc, title, level=1)
            elif title.startswith("Chương"):
                add_heading(doc, title, level=1)
            elif title == "Mục lục":
                i += 1
                while i < len(lines) and lines[i].strip().startswith("- "):
                    i += 1
                continue
            elif title in (
                "Danh mục từ viết tắt",
                "Danh sách hình vẽ",
                "Danh sách bảng biểu",
            ):
                add_heading(doc, title, level=1)
            i += 1
            continue

        if stripped.startswith("## "):
            add_heading(doc, stripped[3:].strip(), level=2)
            i += 1
            continue

        if stripped.startswith("### "):
            add_heading(doc, stripped[4:].strip(), level=3)
            i += 1
            continue

        if stripped.startswith("**") and stripped.endswith("**") and "Bảng" in stripped or "Hình" in stripped:
            p = add_paragraph(doc)
            run = p.add_run(stripped.strip("*"))
            set_run_font(run, bold=True, italic=True)
            i += 1
            continue

        if stripped.startswith("- "):
            item = stripped[2:]
            item = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", item)
            p = add_paragraph(doc, style="List Bullet")
            set_paragraph_format(p, indent_cm=0.5)
            run = p.add_run(item)
            set_run_font(run)
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if m:
            p = add_paragraph(doc)
            set_paragraph_format(p, indent_cm=0.5)
            run = p.add_run(f"{m.group(1)}. {m.group(2)}")
            set_run_font(run)
            i += 1
            continue

        if stripped.startswith("**") and "**:" in stripped:
            p = add_paragraph(doc)
            parts = re.split(r"\*\*", stripped)
            for idx, part in enumerate(parts):
                if not part:
                    continue
                run = p.add_run(part)
                set_run_font(run, bold=(idx % 2 == 1))
            i += 1
            continue

        if stripped:
            clean = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", stripped)
            clean = clean.replace("**", "")
            add_paragraph(doc, clean)
        i += 1

    if table_buf:
        add_table_from_rows(doc, parse_table_block(table_buf))

    doc.save(out_path)
    return out_path


def add_manual_toc(doc):
    """Mục lục dạng văn bản (Word TOC field cần mở bằng MS Word)."""
    entries = [
        "Danh mục từ viết tắt",
        "Danh sách hình vẽ",
        "Danh sách bảng biểu",
        "Chương 1. Mở đầu",
        "Chương 2. Phân tích yêu cầu và use case",
        "Chương 3. Thiết kế tổng quan ứng dụng",
        "Chương 4. Thiết kế chi tiết chức năng cá nhân đảm nhận",
        "Chương 5. Cài đặt và triển khai ứng dụng",
        "Chương 6. Kết quả thực hiện, thử nghiệm và đánh giá",
        "Kết luận",
        "Tài liệu tham khảo",
    ]
    add_heading(doc, "Mục lục", level=1)
    for idx, title in enumerate(entries, 1):
        p = add_paragraph(doc)
        set_paragraph_format(p, indent_cm=0.5)
        run = p.add_run(f"{idx}. {title}")
        set_run_font(run)
    doc.add_page_break()


def main():
    base = Path(__file__).resolve().parent
    md = base / "bao_cao_btl_lap_trinh_mang.md"
    out = base / "Bao_cao_BTL_15trang_font13.docx"
    if len(sys.argv) > 1:
        out = Path(sys.argv[1])
    convert_markdown(md, out)
    print(f"Created: {out}")


if __name__ == "__main__":
    main()
